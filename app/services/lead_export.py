from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from app.models.lead import CompanyInput
from app.models.outreach import OutreachBrief
from app.models.outreach_message import OutreachMessageSet


def _join_list(
    values: list[str] | None,
) -> str:
    """
    Convert a list of strings into a readable
    single-cell representation for CSV/Excel.
    """

    if not values:
        return ""

    return " | ".join(
        value.strip()
        for value in values
        if value and value.strip()
    )


def _get_verification_fields(
    company: CompanyInput,
) -> tuple[
    str,
    bool,
    bool,
    str,
]:
    """
    Extract verification information from the
    company's signals.

    The strongest verification level found across
    signals is returned.
    """

    levels = {
        "EXACT_JOB_VERIFIED": 4,
        "CAREERS_CONFIRMED": 3,
        "EXTERNAL_ONLY": 2,
        "NOT_VERIFIED": 1,
    }

    strongest_level = "NOT_VERIFIED"
    official_careers_found = False
    exact_job_verified = False
    careers_url = ""

    strongest_rank = 0

    for signal in company.signals:

        level = getattr(
            signal,
            "verification_level",
            "NOT_VERIFIED",
        )

        rank = levels.get(
            level,
            1,
        )

        if rank > strongest_rank:
            strongest_rank = rank
            strongest_level = level

        if getattr(
            signal,
            "official_careers_found",
            False,
        ):
            official_careers_found = True

        if getattr(
            signal,
            "exact_job_verified",
            False,
        ):
            exact_job_verified = True

        signal_careers_url = getattr(
            signal,
            "careers_url",
            None,
        )

        if (
            signal_careers_url
            and not careers_url
        ):
            careers_url = signal_careers_url

    return (
        strongest_level,
        official_careers_found,
        exact_job_verified,
        careers_url,
    )


def _get_primary_signal(
    company: CompanyInput,
) -> object | None:
    """
    Return the strongest available signal.

    Signals are ranked primarily by strength.
    """

    if not company.signals:
        return None

    return max(
        company.signals,
        key=lambda signal: (
            getattr(
                signal,
                "strength",
                0,
            ),
            getattr(
                signal,
                "source_trust",
                0,
            ),
        ),
    )


def build_lead_record(
    company: CompanyInput,
    result: object,
    outreach: OutreachBrief,
    messages: OutreachMessageSet,
) -> dict:
    """
    Build one flat, spreadsheet-friendly lead record.

    This function deliberately keeps the export layer
    separate from discovery, scoring, verification,
    and outreach generation.
    """

    (
        verification_level,
        official_careers_found,
        exact_job_verified,
        careers_url,
    ) = _get_verification_fields(
        company
    )

    primary_signal = _get_primary_signal(
        company
    )

    trigger_evidence = outreach.trigger_evidence

    evidence_url = ""

    if primary_signal is not None:
        evidence_url = (
            getattr(
                primary_signal,
                "source_url",
                None,
            )
            or ""
        )

    if not evidence_url:
        for item in trigger_evidence:
            if (
                item.startswith("http://")
                or item.startswith("https://")
            ):
                evidence_url = item
                break

    email_subjects = list(
        getattr(
            messages,
            "cold_email_subjects",
            [],
        )
        or []
    )

    return {
        "Company": company.company_name,

        "Website": company.website or "",

        "Country": company.country or "",

        "Industry": company.industry or "",

        "Employees": (
            company.employee_count
            if company.employee_count is not None
            else ""
        ),

        "Priority": result.priority,

        "Opportunity Score": (
            result.opportunity_score
        ),

        "ICP Score": result.icp_score,

        "Intent Score": result.intent_score,

        "Evidence Quality": (
            result.evidence_quality
        ),

        "Recommended Service": (
            result.recommended_service
        ),

        "Verification Level": (
            verification_level
        ),

        "Official Careers Found": (
            "YES"
            if official_careers_found
            else "NO"
        ),

        "Exact Job Verified": (
            "YES"
            if exact_job_verified
            else "NO"
        ),

        "Careers URL": careers_url,

        "Trigger": outreach.trigger,

        "Trigger Evidence": _join_list(
            trigger_evidence
        ),

        "Evidence URL": evidence_url,

        "Recommended Action": (
            outreach.recommended_action
        ),

        "Decision Maker": (
            outreach.decision_maker_role
        ),

        "ESSEMVEE Angle": (
            outreach.essemvee_angle
        ),

        "LinkedIn Connection": (
            messages.linkedin_connection
        ),

        "LinkedIn Follow-up": (
            messages.linkedin_follow_up
        ),

        "Email Subject 1": (
            email_subjects[0]
            if len(email_subjects) > 0
            else ""
        ),

        "Email Subject 2": (
            email_subjects[1]
            if len(email_subjects) > 1
            else ""
        ),

        "Email Subject 3": (
            email_subjects[2]
            if len(email_subjects) > 2
            else ""
        ),

        "Cold Email": messages.cold_email,

        "Call Opener": messages.call_opener,

        "Discovery Question": (
            messages.discovery_question
        ),

        "CTA": messages.cta,

        "Personalization Points": _join_list(
            messages.personalization_points
        ),

        "Evidence Gaps": _join_list(
            outreach.evidence_gaps
        ),

        "Do Not Claim": _join_list(
            messages.do_not_claim
        ),
    }


def export_csv(
    records: Iterable[dict],
    output_path: str | Path,
) -> Path:
    """
    Export lead records to CSV.
    """

    records = list(records)

    output = Path(
        output_path
    )

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not records:
        raise ValueError(
            "Cannot export an empty lead set."
        )

    fieldnames = list(
        records[0].keys()
    )

    with output.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
        )

        writer.writeheader()

        writer.writerows(
            records
        )

    return output


def export_excel(
    records: Iterable[dict],
    output_path: str | Path,
) -> Path:
    """
    Export lead records to an Excel workbook.

    Uses openpyxl so the resulting file is directly
    usable in Excel, Google Sheets, or CRM import
    workflows.
    """

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    records = list(records)

    output = Path(
        output_path
    )

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not records:
        raise ValueError(
            "Cannot export an empty lead set."
        )

    fieldnames = list(
        records[0].keys()
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Leads"

    # Header
    for column_index, field in enumerate(
        fieldnames,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=field,
        )

        cell.font = Font(
            bold=True
        )

    # Data
    for row_index, record in enumerate(
        records,
        start=2,
    ):

        for column_index, field in enumerate(
            fieldnames,
            start=1,
        ):

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=record.get(
                    field,
                    "",
                ),
            )

    # Freeze header
    worksheet.freeze_panes = "A2"

    # Enable filtering
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Reasonable column widths
    for column_index, field in enumerate(
        fieldnames,
        start=1,
    ):

        max_length = len(field)

        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):

            value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is not None:

                max_length = max(
                    max_length,
                    min(
                        len(str(value)),
                        60,
                    ),
                )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = min(
            max_length + 2,
            60,
        )

    workbook.save(
        output
    )

    return output